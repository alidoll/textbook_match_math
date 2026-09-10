"""基准库 xlsx 解析（表头、行字段、课名拆分）。"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

from openpyxl import Workbook, load_workbook

SKIP_SHEETS = frozenset({"小学科学方案", "初中科学内容生产"})

CN_GRADE = {
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
}


@dataclass(frozen=True)
class BenchmarkLessonRow:
    sheet: str
    version_label: str
    grade: int
    semester: str  # 上 | 下
    unit_title: str
    lesson_raw: str
    lesson_no: str
    lesson_name: str
    old_course_id: str
    page_count: int | None
    import_batch: str | None


def cell_str(value: Any) -> str:
    if value is None or (isinstance(value, float) and str(value) == "nan"):
        return ""
    s = str(value).strip()
    return "" if s in {"None", "nan"} else s


def build_header_index(header_row: tuple[Any, ...]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        key = cell_str(cell)
        if key:
            out[key] = i
    return out


def row_get(row: tuple[Any, ...], colmap: dict[str, int], *names: str) -> Any:
    for name in names:
        if name in colmap:
            idx = colmap[name]
            if idx < len(row):
                return row[idx]
    return None


def parse_grade(value: Any) -> int | None:
    s = cell_str(value).replace("年级", "").replace(" ", "")
    if not s:
        return None
    if s.isdigit():
        return int(s)
    for ch, num in CN_GRADE.items():
        if s.startswith(ch):
            return num
    return None


def parse_semester(value: Any) -> str:
    t = cell_str(value)
    if "下" in t:
        return "下"
    return "上"


def parse_lesson_cell(raw: Any) -> tuple[str, str]:
    s = cell_str(raw)
    if not s:
        return "", ""
    # 兼容「1 课名」「1.课名」「1．课名」「1、课名」
    m = re.match(r"^(\d+)\s*[.．、]?\s*(.*)$", s)
    if m:
        no, name = m.group(1), m.group(2).strip()
        name = name.lstrip(".．、").strip()
        return no, name or s
    return "", s.lstrip(".．、").strip() or s


def lesson_no_sort_key(lesson_raw: str) -> tuple[int, int, str]:
    """目录/课时按节号排序；无节号或单元小结排在数字课后。"""
    s = cell_str(lesson_raw)
    if "单元小结" in s:
        return (1, 9999, s)
    m = re.match(r"^(\d+)", s)
    if m:
        return (0, int(m.group(1)), s)
    return (2, 0, s)


def sort_catalog_rows(
    catalog: list[dict[str, Any]],
    *,
    subject: str | None = None,
) -> list[dict[str, Any]]:
    """按单元归位后再排课。

    科学/数学等所见即所得：单元内保持识别顺序（勿用可能标错的课号重排）。
    化学/语文：单元内再按节号排（OCR 多列常乱序）。
    """
    from .unit_title import (
        sort_catalog_by_unit_and_lesson,
        sort_catalog_by_unit_preserving_order,
    )

    sub = (subject or "").strip()
    # 与 lesson_filters.catalog_filter_mode 对齐，避免 parsers→services 循环依赖
    if sub in ("语文", "Test", "化学"):
        return sort_catalog_by_unit_and_lesson(catalog)
    return sort_catalog_by_unit_preserving_order(catalog)


def format_grade_cell(grade: int) -> str:
    return f"{grade}年级"


def format_semester_cell(semester: str) -> str:
    s = (semester or "").strip()
    if s in {"上", "上册"}:
        return "上学期"
    if "下" in s or s in {"下册"}:
        return "下学期"
    return s


def _semester_matches_cell(semester: str, cell: str) -> bool:
    want = parse_semester(semester)
    got = parse_semester(cell)
    return want == got


def append_catalog_to_benchmark_xlsx(
    xlsx_path,
    sheet_name: str,
    *,
    version_label: str,
    grade: int,
    semester: str,
    catalog: list[dict[str, Any]],
    replace_existing: bool = True,
    import_batch: str | None = None,
) -> tuple[int, Path, tuple[int, int] | None]:
    """
    将 PDF 目录行写入新课标基准 xlsx（版本/年级/学期/单元/节；不写页数、完成时间）。
    返回 (写入行数, xlsx 路径, 起止行号 1-based 或 None)。
    """
    path = Path(xlsx_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    headers = ["版本", "年级", "学期", "单元", "节", "课件id", "页数", "完成时间"]
    try:
        if path.is_file():
            wb = load_workbook(path)
        else:
            wb = Workbook()
            if wb.sheetnames:
                wb.remove(wb.active)
    except (PermissionError, OSError) as exc:
        if isinstance(exc, OSError) and exc.errno != 13:
            raise
        raise PermissionError(
            f"无法打开基准库文件（请先关闭 Excel/WPS 中打开的该文件）：{path}"
        ) from exc

    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
    else:
        ws = wb[sheet_name]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row or not any(header_row):
        ws.delete_rows(1, 1)
        ws.insert_rows(1)
        for i, name in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=name)
        header_row = tuple(headers)
    colmap = build_header_index(tuple(header_row))

    grade_cell = format_grade_cell(grade)
    semester_cell = format_semester_cell(semester)
    _ = import_batch  # 仅用于调用方标记批次，不写入表格

    def _cell(row_idx: int, col_name: str) -> Any:
        if col_name not in colmap:
            return None
        return ws.cell(row=row_idx, column=colmap[col_name] + 1).value

    if replace_existing and ws.max_row > 1:
        for r in range(ws.max_row, 1, -1):
            ver = cell_str(_cell(r, "版本"))
            gr = parse_grade(_cell(r, "年级"))
            sem_raw = cell_str(_cell(r, "学期"))
            if (
                ver == version_label
                and gr == grade
                and _semester_matches_cell(semester, sem_raw)
            ):
                ws.delete_rows(r, 1)

    written = 0
    row_span: tuple[int, int] | None = None
    start_row = ws.max_row + 1
    for item in catalog:
        lesson = cell_str(item.get("lesson"))
        unit = cell_str(item.get("unit"))
        if not lesson or not unit:
            continue
        ws.append(
            [
                version_label,
                grade_cell,
                semester_cell,
                unit,
                lesson,
                "",
            ]
        )
        written += 1
    if written:
        row_span = (start_row, ws.max_row)

    try:
        wb.save(path)
    except (PermissionError, OSError) as exc:
        wb.close()
        if isinstance(exc, OSError) and exc.errno != 13:
            raise
        raise PermissionError(
            f"无法保存基准库（请先关闭 Excel/WPS 中打开的该文件）：{path}"
        ) from exc
    wb.close()
    return written, path, row_span


def find_volume_catalog_row_span(
    xlsx_path,
    sheet_name: str,
    *,
    version_label: str,
    grade: int,
    semester: str,
) -> tuple[int, int] | None:
    """按版本/年级/学期查找基准库中对应目录行的起止行号。"""
    path = Path(xlsx_path)
    if not path.is_file():
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            return None
        colmap = build_header_index(tuple(header))
        rows: list[int] = []
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            ver = cell_str(row_get(row, colmap, "版本"))
            gr = parse_grade(row_get(row, colmap, "年级"))
            sem_raw = cell_str(row_get(row, colmap, "学期"))
            if (
                ver == version_label
                and gr == grade
                and _semester_matches_cell(semester, sem_raw)
            ):
                rows.append(r)
        if not rows:
            return None
        return rows[0], rows[-1]
    finally:
        wb.close()


def find_import_batch_row_span(
    xlsx_path,
    sheet_name: str,
    import_batch: str,
) -> tuple[int, int] | None:
    """已弃用：完成时间列不再写入 pdf 批次标记。"""
    _ = import_batch
    return None


def merge_misgraded_lesson_rows(
    primary: list[BenchmarkLessonRow],
    candidates: list[BenchmarkLessonRow],
    *,
    grade: int,
) -> list[BenchmarkLessonRow]:
    """
    补入基准库里「年级填错」但节号落本册范围内的行。
    例：湘科二下第 7 课「植物和我们」误标为三年级，与二下 4–6 课同单元。
    """
    primary_nos = sorted(
        int(r.lesson_no) for r in primary if str(r.lesson_no).isdigit()
    )
    if not primary_nos:
        return list(primary)

    lo, hi = primary_nos[0], primary_nos[-1]
    have = {int(r.lesson_no) for r in primary if str(r.lesson_no).isdigit()}
    primary_units = {r.unit_title for r in primary}

    extras: list[BenchmarkLessonRow] = []
    for row in candidates:
        if row.grade == grade:
            continue
        if not str(row.lesson_no).isdigit():
            continue
        n = int(row.lesson_no)
        if n in have or n < lo or n > hi:
            continue
        if row.unit_title not in primary_units:
            continue
        extras.append(row)
        have.add(n)

    merged = list(primary) + extras
    merged.sort(key=lambda r: lesson_no_sort_key(r.lesson_raw or r.lesson_name))
    return merged


def parse_page_count(value: Any) -> int | None:
    s = cell_str(value)
    if not s:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def iter_sheet_rows(xlsx_path, sheet_name: str) -> Iterator[BenchmarkLessonRow]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise FileNotFoundError(f"基准库缺少工作表：{sheet_name}")
        ws = wb[sheet_name]
        row_iter = ws.iter_rows(values_only=True)
        header = next(row_iter, None)
        if not header:
            return
        colmap = build_header_index(tuple(header))
        if "节" not in colmap:
            raise ValueError(f"工作表 {sheet_name} 缺少「节」列")

        for row in row_iter:
            if not row or not row_get(row, colmap, "节"):
                continue
            grade_raw = row_get(row, colmap, "年级")
            if cell_str(grade_raw) in {"", "年级"}:
                continue
            grade = parse_grade(grade_raw)
            if grade is None:
                continue
            lesson_no, lesson_name = parse_lesson_cell(row_get(row, colmap, "节"))
            yield BenchmarkLessonRow(
                sheet=sheet_name,
                version_label=cell_str(row_get(row, colmap, "版本")),
                grade=grade,
                semester=parse_semester(row_get(row, colmap, "学期")),
                unit_title=cell_str(row_get(row, colmap, "单元")),
                lesson_raw=cell_str(row_get(row, colmap, "节")),
                lesson_no=lesson_no,
                lesson_name=lesson_name,
                old_course_id=cell_str(
                    row_get(row, colmap, "课件id", "课件ID", "课件唯一ID")
                ),
                page_count=parse_page_count(row_get(row, colmap, "页数")),
                import_batch=cell_str(row_get(row, colmap, "完成时间")) or None,
            )
    finally:
        wb.close()


def list_workbook_sheets(xlsx_path) -> list[str]:
    wb = load_workbook(xlsx_path, read_only=True)
    try:
        return [s for s in wb.sheetnames if s not in SKIP_SHEETS]
    finally:
        wb.close()
