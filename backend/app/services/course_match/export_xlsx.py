"""新库粗分结果导出 Excel（单册 / 本版多册）。"""
from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ...models import Lesson
from ..new_library.lesson_validation import lesson_validation_for_lesson
from ..old_library.edition_registry import get_edition, grade_term_pairs_for_edition
from ..old_library.volume_codes import make_volume_code, normalize_term
from .service import get_course_match_result

_INVALID_FS = re.compile(r'[\\/:*?"<>|]+')
_INVALID_SHEET = re.compile(r"[\\/*?:\[\]]+")

_HEADERS = [
    "序号",
    "新课单元",
    "新课节号",
    "新课名称",
    "新课页码起",
    "新课页码止",
    "主对照档位",
    "参考相似度",
    "主对照旧课",
    "旧课册次",
    "旧课课件ID",
    "溯源(rank2)",
    "备选(rank3)",
    "备选(rank4)",
    "系统建块建议",
    "验收状态",
    "预判断结论",
    "备注",
    "新课 lesson_uid",
    "册次代码",
]

_UNMATCHED_TIERS = frozenset({"none", "fully_new", "no_match", ""})


def _pct(score: float | None) -> str:
    if score is None:
        return ""
    return f"{score * 100:.1f}%"


def _candidate_summary(candidates: list[dict], rank: int) -> str:
    for c in candidates or []:
        if int(c.get("rank") or 0) == rank:
            tier = c.get("match_label") or c.get("match_tier") or ""
            hint = (c.get("old_lesson_hint") or "").replace("\n", " ")
            score = _pct(c.get("similarity_score"))
            vol = c.get("old_volume_code") or ""
            parts = [p for p in [tier, score, vol, hint] if p]
            return " | ".join(parts)
    return ""


def _is_matched_row(row: dict[str, Any]) -> bool:
    tier = str(row.get("match_tier") or "").strip().lower()
    if tier in _UNMATCHED_TIERS:
        return False
    if row.get("old_lesson_hint") or row.get("old_courseware_id"):
        return True
    cands = row.get("candidates") or []
    if cands and cands[0].get("old_lesson_id"):
        return True
    return tier not in _UNMATCHED_TIERS and bool(tier)


def _filter_rows(
    rows: list[dict[str, Any]],
    *,
    include_matched: bool,
    include_unmatched: bool,
) -> list[dict[str, Any]]:
    if include_matched and include_unmatched:
        return list(rows)
    out: list[dict[str, Any]] = []
    for row in rows:
        matched = _is_matched_row(row)
        if matched and include_matched:
            out.append(row)
        elif (not matched) and include_unmatched:
            out.append(row)
    return out


def _safe_filename(name: str, *, fallback: str = "粗分结果") -> str:
    s = _INVALID_FS.sub("_", (name or "").strip()) or fallback
    if not s.lower().endswith(".xlsx"):
        s = f"{s}.xlsx"
    return s


def _safe_sheet_name(name: str, used: set[str]) -> str:
    base = _INVALID_SHEET.sub("", (name or "").strip()) or "粗分"
    base = base[:31]
    candidate = base
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = f"{base[: max(1, 31 - len(suffix))]}{suffix}"
        i += 1
    used.add(candidate)
    return candidate


def _append_detail_sheet(
    wb: Workbook,
    *,
    title: str,
    volume_code: str,
    rows: list[dict[str, Any]],
    used_names: set[str],
) -> None:
    ws = wb.create_sheet(_safe_sheet_name(title, used_names))
    ws.append(_HEADERS)
    for col in range(1, len(_HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    uids = [r["new_lesson_uid"] for r in rows if r.get("new_lesson_uid")]
    lessons = {
        les.lesson_uid: les
        for les in Lesson.query.filter(Lesson.lesson_uid.in_(uids)).all()
    } if uids else {}

    flag_fill = PatternFill(fill_type="solid", fgColor="FFF7D6")
    for i, row in enumerate(rows, start=1):
        uid = row.get("new_lesson_uid") or ""
        les = lessons.get(uid)
        cands = row.get("candidates") or []
        validation = lesson_validation_for_lesson(les.id) if les else {}
        ws.append(
            [
                i,
                row.get("new_unit_title") or "",
                row.get("new_lesson_no") or "",
                row.get("new_lesson_name") or "",
                les.page_start if les else "",
                les.page_end if les else "",
                row.get("match_label") or row.get("match_tier") or "",
                _pct(row.get("similarity_score")),
                row.get("old_lesson_hint") or "",
                (cands[0].get("old_volume_code") if cands else "") or "",
                row.get("old_courseware_id") or "",
                _candidate_summary(cands, 2),
                _candidate_summary(cands, 3),
                _candidate_summary(cands, 4),
                row.get("build_hint_label") or "",
                validation.get("validation_label") or "",
                validation.get("prescan_summary")
                or validation.get("prescan_agreement_label")
                or "",
                "",
                uid,
                volume_code,
            ]
        )
        if validation.get("needs_review"):
            excel_row = ws.max_row
            for c in range(1, len(_HEADERS) + 1):
                ws.cell(row=excel_row, column=c).fill = flag_fill

    widths = [6, 18, 8, 22, 8, 8, 10, 10, 48, 12, 14, 40, 40, 40, 10, 14, 36, 20, 28, 16]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(_HEADERS) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(
                wrap_text=True, vertical="top"
            )


def _volume_sheet_title(volume_code: str, grade: int | None = None, term: str | None = None) -> str:
    if grade is not None and term:
        term_l = "上册" if normalize_term(term) == "上" else "下册"
        return f"{grade}年级{term_l}"
    return volume_code


def build_volume_course_match_xlsx(
    *,
    volume_code: str,
    include_matched: bool = True,
    include_unmatched: bool = True,
) -> tuple[bytes, str]:
    """导出单册粗分 Excel。返回 (bytes, filename)。"""
    if not include_matched and not include_unmatched:
        raise ValueError("请至少勾选「已匹配」或「未匹配」")
    data = get_course_match_result(volume_code=volume_code)
    rows = _filter_rows(
        list(data.get("matches_by_lesson") or []),
        include_matched=include_matched,
        include_unmatched=include_unmatched,
    )
    if not data.get("job_id"):
        raise ValueError(f"{volume_code} 尚无粗分结果，请先运行粗分")
    if not rows:
        raise ValueError(f"{volume_code} 按当前筛选无课时可导出")

    wb = Workbook()
    ws_meta = wb.active
    ws_meta.title = "说明"
    summary = data.get("summary") or {}
    filter_label = []
    if include_matched:
        filter_label.append("已匹配")
    if include_unmatched:
        filter_label.append("未匹配")
    for row in [
        ["册次代码", volume_code],
        ["粗分任务 ID", data.get("job_id") or ""],
        ["完成时间", data.get("finished_at") or ""],
        ["筛选", "、".join(filter_label)],
        ["导出课时数", len(rows)],
        ["完全同名", summary.get("exact_match", "")],
        ["高相似", summary.get("high_similarity", "")],
        ["溯源", summary.get("traceability", "")],
        ["无匹配", summary.get("no_match", "")],
        ["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]:
        ws_meta.append(row)
    ws_meta.column_dimensions["A"].width = 14
    ws_meta.column_dimensions["B"].width = 72

    used: set[str] = {"说明"}
    _append_detail_sheet(
        wb,
        title="粗分明细",
        volume_code=volume_code,
        rows=rows,
        used_names=used,
    )
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), _safe_filename(f"粗分_{volume_code}")


def list_edition_course_match_export_targets(edition_id: str) -> list[dict[str, Any]]:
    """本版可导出粗分的册次（有粗分任务即可）。"""
    edition = get_edition(edition_id)
    rows: list[dict[str, Any]] = []
    for grade, term in grade_term_pairs_for_edition(edition):
        term_key = normalize_term(term)
        code = make_volume_code(edition, grade=grade, term=term_key, book_type="new")
        try:
            data = get_course_match_result(volume_code=code)
        except ValueError:
            rows.append(
                {
                    "volume_code": code,
                    "grade": grade,
                    "term": term_key,
                    "label": _volume_sheet_title(code, grade, term_key),
                    "has_course_match": False,
                    "lesson_count": 0,
                    "matched_count": 0,
                    "unmatched_count": 0,
                    "status": None,
                }
            )
            continue
        has_job = bool(data.get("job_id"))
        matches = list(data.get("matches_by_lesson") or [])
        matched_n = sum(1 for r in matches if _is_matched_row(r))
        unmatched_n = len(matches) - matched_n
        rows.append(
            {
                "volume_code": code,
                "grade": grade,
                "term": term_key,
                "label": _volume_sheet_title(code, grade, term_key),
                "has_course_match": has_job,
                "lesson_count": len(matches),
                "matched_count": matched_n,
                "unmatched_count": unmatched_n,
                "status": data.get("status"),
            }
        )
    return rows


def build_edition_course_match_xlsx(
    *,
    edition_id: str,
    volume_codes: Iterable[str] | None = None,
    include_matched: bool = True,
    include_unmatched: bool = True,
) -> tuple[bytes, str]:
    """导出本版多册粗分 Excel（每册一表）。"""
    if not include_matched and not include_unmatched:
        raise ValueError("请至少勾选「已匹配」或「未匹配」")
    edition = get_edition(edition_id)
    wanted = {str(c).strip() for c in (volume_codes or []) if str(c).strip()}
    targets = list_edition_course_match_export_targets(edition.edition_id)
    if wanted:
        targets = [t for t in targets if t["volume_code"] in wanted]
        missing = wanted - {t["volume_code"] for t in targets}
        if missing:
            raise ValueError(f"册次不在本版：{', '.join(sorted(missing))}")
    exportable = [t for t in targets if t["has_course_match"]]
    if not exportable:
        raise ValueError("所选册次尚无粗分结果，请先运行粗分")

    wb = Workbook()
    ws_meta = wb.active
    ws_meta.title = "说明"
    filter_label = []
    if include_matched:
        filter_label.append("已匹配")
    if include_unmatched:
        filter_label.append("未匹配")
    ws_meta.append(["版本", edition.label])
    ws_meta.append(["版本 ID", edition.edition_id])
    ws_meta.append(["筛选", "、".join(filter_label)])
    ws_meta.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws_meta.append([])
    ws_meta.append(["册次", "已匹配", "未匹配", "导出行数", "状态"])

    used: set[str] = {"说明"}
    total_rows = 0
    for t in exportable:
        data = get_course_match_result(volume_code=t["volume_code"])
        rows = _filter_rows(
            list(data.get("matches_by_lesson") or []),
            include_matched=include_matched,
            include_unmatched=include_unmatched,
        )
        ws_meta.append(
            [
                t["label"],
                t["matched_count"],
                t["unmatched_count"],
                len(rows),
                data.get("status") or "",
            ]
        )
        if not rows:
            continue
        _append_detail_sheet(
            wb,
            title=t["label"],
            volume_code=t["volume_code"],
            rows=rows,
            used_names=used,
        )
        total_rows += len(rows)

    if total_rows <= 0:
        raise ValueError("按当前筛选无课时可导出（请勾选已匹配/未匹配，或换有结果的册）")

    ws_meta.column_dimensions["A"].width = 16
    ws_meta.column_dimensions["B"].width = 12
    ws_meta.column_dimensions["C"].width = 12
    ws_meta.column_dimensions["D"].width = 12
    ws_meta.column_dimensions["E"].width = 12

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), _safe_filename(f"{edition.label}_本版粗分")
