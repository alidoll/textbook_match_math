"""教材对比页：文字比对 + 图片比对导出为双工作表 Excel。"""
from __future__ import annotations

from io import BytesIO
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font

from ...models import Volume
from .atom_compare import compare_cached_page_text
from .page_image_compare import compare_cached_page_images


def _ops_summary(ops: list[dict] | None) -> str:
    if not ops:
        return ""
    parts: list[str] = []
    for op in ops[:20]:
        kind = (op.get("kind") or "").strip()
        old = (op.get("old") or "").strip()
        new = (op.get("new") or "").strip()
        if kind == "文字删除" and old:
            parts.append(f"删「{old}」")
        elif kind == "文字新增" and new:
            parts.append(f"增「{new}」")
        elif kind == "文字改写":
            parts.append(f"改「{old}」→「{new}」")
        elif kind:
            bits = [kind]
            if old:
                bits.append(old)
            if new:
                bits.append(f"→{new}")
            parts.append(" ".join(bits))
    return "；".join(parts)


def _image_rows(compare: dict[str, Any] | None) -> list[dict[str, Any]]:
    """与前端 normalizeImageBlockRows 对齐。"""
    if not compare:
        return []
    layout = compare.get("layout") or {}
    caps_by_key: dict[str, dict] = {}
    for r in (compare.get("captions") or {}).get("rows") or []:
        caps_by_key[f"{r.get('old_atom_id')}|{r.get('new_atom_id')}"] = r
    vis_by_key: dict[str, dict] = {}
    for r in (compare.get("visuals") or {}).get("rows") or []:
        vis_by_key[f"{r.get('old_atom_id')}|{r.get('new_atom_id')}"] = r

    rows: list[dict[str, Any]] = []
    for p in layout.get("pairs") or []:
        key = f"{p.get('old_atom_id')}|{p.get('new_atom_id')}"
        cap = caps_by_key.get(key) or {}
        vis = vis_by_key.get(key) or {}
        rows.append(
            {
                "kind": "pair",
                "old_atom_id": p.get("old_atom_id"),
                "new_atom_id": p.get("new_atom_id"),
                "layout_change": p.get("change") or "",
                "old_theme": cap.get("old_theme") or cap.get("old_label") or "",
                "new_theme": cap.get("new_theme") or cap.get("new_label") or "",
                "theme_change": cap.get("theme_change") or "",
                "old_caption": cap.get("old_caption") or "",
                "new_caption": cap.get("new_caption") or "",
                "caption_change": cap.get("caption_change") or cap.get("change") or "",
                "visual_grade": vis.get("grade") or "",
                "visual_similarity": vis.get("similarity"),
            }
        )
    for u in layout.get("unmatched_old") or []:
        rows.append(
            {
                "kind": "old_only",
                "old_atom_id": u.get("atom_id"),
                "new_atom_id": None,
                "layout_change": "旧有新无",
                "old_theme": "",
                "new_theme": "",
                "theme_change": "",
                "old_caption": "",
                "new_caption": "",
                "caption_change": "",
                "visual_grade": "",
                "visual_similarity": None,
            }
        )
    for u in layout.get("unmatched_new") or []:
        rows.append(
            {
                "kind": "new_only",
                "old_atom_id": None,
                "new_atom_id": u.get("atom_id"),
                "layout_change": "新有旧无",
                "old_theme": "",
                "new_theme": "",
                "theme_change": "",
                "old_caption": "",
                "new_caption": "",
                "caption_change": "",
                "visual_grade": "",
                "visual_similarity": None,
            }
        )
    return rows


def _image_change_label(row: dict[str, Any]) -> str:
    if row.get("kind") == "old_only":
        return "旧有新无"
    if row.get("kind") == "new_only":
        return "新有旧无"
    issues: list[str] = []
    lc = row.get("layout_change") or ""
    if lc and lc != "版面对齐":
        issues.append(lc)
    if row.get("theme_change") == "主题改写":
        issues.append("主题改写")
    cc = row.get("caption_change") or ""
    if cc == "图注改写":
        issues.append("图注改写")
    vg = row.get("visual_grade") or ""
    if vg and vg != "画面接近":
        issues.append(vg)
    return issues[0] if issues else "没变化"


def _style_header(ws) -> None:
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, min_width: int = 10, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = col[0].column_letter
        width = min_width
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(val) + 2))
        ws.column_dimensions[letter].width = width


def build_compare_export_xlsx(
    *,
    old_vol: Volume,
    new_vol: Volume,
    old_page: int,
    new_page: int,
    new_pdf_source: str = "full",
    preview_blob_id: str | None = None,
) -> tuple[bytes, str]:
    """
    生成双工作表 xlsx：文字对比 / 图片对比。
    返回 (xlsx_bytes, download_name)。
    """
    text_cmp: dict[str, Any] | None = None
    image_cmp: dict[str, Any] | None = None
    text_err = ""
    image_err = ""

    try:
        text_out = compare_cached_page_text(
            old_vol=old_vol,
            new_vol=new_vol,
            old_page=old_page,
            new_page=new_page,
            new_pdf_source=new_pdf_source,
            preview_blob_id=preview_blob_id,
            force=False,
        )
        text_cmp = text_out.get("compare")
    except ValueError as exc:
        text_err = str(exc)

    try:
        image_out = compare_cached_page_images(
            old_vol=old_vol,
            new_vol=new_vol,
            old_page=old_page,
            new_page=new_page,
            new_pdf_source=new_pdf_source,
            preview_blob_id=preview_blob_id,
        )
        image_cmp = image_out.get("compare")
    except ValueError as exc:
        image_err = str(exc)

    if not text_cmp and not image_cmp:
        raise ValueError(
            text_err or image_err or "尚无比对结果，请先完成 ② 文字比对 与/或 ④ 图片比对"
        )

    wb = openpyxl.Workbook()

    # —— 文字对比 ——
    ws_text = wb.active
    ws_text.title = "文字对比"
    ws_text.append(
        [
            "旧原子ID",
            "新原子ID",
            "旧教材文本",
            "新教材文本",
            "变化说明",
            "相似度",
            "差异明细",
        ]
    )
    if text_cmp:
        ws_text.append(
            [
                "（摘要）",
                "",
                f"结论：{text_cmp.get('verdict') or '—'}",
                f"相似度：{text_cmp.get('similarity') if text_cmp.get('similarity') is not None else '—'}",
                f"有差异块：{(text_cmp.get('summary') or {}).get('changed_blocks', '—')}/{(text_cmp.get('summary') or {}).get('block_count', '—')}",
                "",
                f"旧PDF p{old_page} ↔ 新PDF p{new_page}",
            ]
        )
        for row in text_cmp.get("block_rows") or []:
            sim = row.get("similarity")
            sim_s = f"{float(sim):.3f}" if sim is not None else ""
            detail = (row.get("change_summary") or "").strip() or _ops_summary(row.get("ops"))
            ws_text.append(
                [
                    row.get("old_atom_id") or "",
                    row.get("new_atom_id") or "",
                    row.get("old_text") or "",
                    row.get("new_text") or "",
                    row.get("change") or "",
                    sim_s,
                    detail,
                ]
            )
    else:
        ws_text.append(["（无文字比对）", "", text_err or "请先完成 ② 文字比对", "", "", "", ""])

    _style_header(ws_text)
    _autosize(ws_text)
    for row in ws_text.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # —— 图片对比 ——
    ws_img = wb.create_sheet("图片对比")
    ws_img.append(
        [
            "旧原子ID",
            "新原子ID",
            "变化说明",
            "版面变化",
            "旧图义",
            "新图义",
            "图义变化",
            "旧图注",
            "新图注",
            "图注变化",
            "画面等级",
            "画面相似度",
        ]
    )
    if image_cmp:
        ws_img.append(
            [
                "（摘要）",
                "",
                image_cmp.get("verdict") or "—",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                f"旧PDF p{old_page} ↔ 新PDF p{new_page}",
            ]
        )
        for row in _image_rows(image_cmp):
            sim = row.get("visual_similarity")
            sim_s = f"{float(sim):.3f}" if sim is not None else ""
            ws_img.append(
                [
                    row.get("old_atom_id") or "",
                    row.get("new_atom_id") or "",
                    _image_change_label(row),
                    row.get("layout_change") or "",
                    row.get("old_theme") or "",
                    row.get("new_theme") or "",
                    row.get("theme_change") or "",
                    row.get("old_caption") or "",
                    row.get("new_caption") or "",
                    row.get("caption_change") or "",
                    row.get("visual_grade") or "",
                    sim_s,
                ]
            )
    else:
        ws_img.append(
            ["（无图片比对）", "", image_err or "请先完成 ④ 图片比对", "", "", "", "", "", "", "", "", ""]
        )

    _style_header(ws_img)
    _autosize(ws_img)
    for row in ws_img.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    old_code = getattr(old_vol, "volume_code", None) or "old"
    new_code = getattr(new_vol, "volume_code", None) or "new"
    name = f"{old_code}_{new_code}_p{old_page}-{new_page}_compare.xlsx"
    return buf.getvalue(), name
