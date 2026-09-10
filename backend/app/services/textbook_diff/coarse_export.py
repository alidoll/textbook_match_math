"""本册建设：课时粗分结果导出 Excel。"""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ...models import Volume
from ..old_library.edition_registry import GRADE_LABELS
from .volumes import get_diff_volume_by_code
from .workbook import list_stored_pairs

_INVALID_FS_CHARS = re.compile(r'[\\/:*?"<>|]+')
_INVALID_SHEET_CHARS = re.compile(r'[\\/*?:\[\]]+')

_HEADER_FILL = PatternFill("solid", fgColor="BE185D")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=14, color="9D174D")
_META_FONT = Font(size=10, color="64748B")
_CELL_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
_WRAP = Alignment(wrap_text=True, vertical="center")
_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

_MATCH_METHOD_LABEL = {
    "name_exact": "完全匹配",
    "name_similar": "相似匹配",
    "name_edition": "跨年级匹配",
    "name_llm": "大模型匹配",
    "none": "未匹配",
    "title_anchor": "课名匹配",
    "lesson_no": "课号匹配",
    "index": "顺序对齐",
    "manual": "手工改对",
    "footer": "印刷页码",
    "position": "页序估计",
    "content_weak": "内容弱匹配",
}

_HEADERS = [
    "序号",
    "出版社",
    "新教材课程名称",
    "新·年级",
    "新·册次",
    "新·单元",
    "新·课时",
    "匹配方式",
    "旧教材课程名称",
    "旧·年级",
    "旧·册次",
    "旧·单元",
    "旧·课时",
    "课件id",
]

# 居中列：序号、出版社、年级、册次、匹配方式、课件id
_CENTER_COLS = frozenset({1, 2, 4, 5, 8, 10, 11, 14})


def _grade_label(grade: int | None) -> str:
    if grade is None:
        return ""
    try:
        g = int(grade)
    except (TypeError, ValueError):
        return str(grade)
    return GRADE_LABELS.get(g, f"{g}年级")


def _semester_label(semester: str | None) -> str:
    s = (semester or "").strip()
    if not s:
        return ""
    if s in ("上", "shang", "S"):
        return "上册"
    if s in ("下", "xia", "X"):
        return "下册"
    if s.endswith("册"):
        return s
    return s


def _vol_fields(vol: Volume | None) -> dict[str, str]:
    if vol is None:
        return {"publisher": "", "grade": "", "volume": ""}
    return {
        "publisher": str(vol.edition or "").strip(),
        "grade": _grade_label(getattr(vol, "grade", None)),
        "volume": _semester_label(getattr(vol, "semester", None)),
    }


def _export_filename(publisher: str, *, fallback: str = "粗分结果") -> str:
    """文件名：出版社名称.xlsx。"""
    raw = _INVALID_FS_CHARS.sub("_", (publisher or "").strip()) or fallback
    return f"{raw}.xlsx"


def _export_sheet_name(grade: str, volume: str, *, fallback: str = "粗分") -> str:
    """工作表名：新教材年级 + 册次（如「四年级上册」）。"""
    raw = f"{(grade or '').strip()}{(volume or '').strip()}".strip() or fallback
    name = _INVALID_SHEET_CHARS.sub("_", raw)[:31]
    return name or fallback


def _lesson_name(brief: dict[str, Any] | None) -> str:
    if not brief:
        return ""
    return str(
        brief.get("lesson_label")
        or brief.get("option_label")
        or brief.get("lesson_name")
        or ""
    ).strip()


def _lesson_unit(brief: dict[str, Any] | None) -> str:
    if not brief:
        return ""
    return str(brief.get("unit_title") or "").strip()


def _lesson_hour(brief: dict[str, Any] | None) -> str:
    """课时列：课号 + 课名（不含单元）。"""
    if not brief:
        return ""
    label = str(brief.get("lesson_label") or "").strip()
    if label:
        return label
    no = str(brief.get("lesson_no") or "").strip()
    name = str(brief.get("lesson_name") or "").strip()
    if no and name:
        return f"{no} {name}".strip()
    return name or no


def _lesson_courseware_id(brief: dict[str, Any] | None) -> str:
    """旧课课件 id（复用定位用）。"""
    if not brief:
        return ""
    return str(brief.get("old_course_id") or "").strip()


def _resolve_old_volume(
    *,
    brief: dict[str, Any] | None,
    fallback: Volume | None,
    volumes_by_code: dict[str, Volume],
) -> Volume | None:
    if not brief:
        return None
    code = str(brief.get("volume_code") or "").strip()
    if code and code in volumes_by_code:
        return volumes_by_code[code]
    if code:
        try:
            return get_diff_volume_by_code(code)
        except ValueError:
            # 小科旧课可能落在旧库册次码上，非 diff 前缀
            from ...models import Volume as VolModel

            row = VolModel.query.filter_by(volume_code=code).first()
            if row:
                return row
    return fallback


def _is_matched_item(it: dict[str, Any]) -> bool:
    method = str(it.get("match_method") or "").strip().lower()
    if method in ("none", "unmatched", "no_match"):
        return False
    return bool(it.get("old"))


def _filter_items(
    items: list[dict[str, Any]],
    *,
    include_matched: bool,
    include_unmatched: bool,
) -> list[dict[str, Any]]:
    if include_matched and include_unmatched:
        return list(items)
    out: list[dict[str, Any]] = []
    for it in items:
        matched = _is_matched_item(it)
        if matched and include_matched:
            out.append(it)
        elif (not matched) and include_unmatched:
            out.append(it)
    return out


def _collect_volumes_by_code(
    items: list[dict[str, Any]],
    *,
    old_vol: Volume,
    new_vol: Volume,
) -> dict[str, Volume]:
    volumes_by_code: dict[str, Volume] = {
        old_vol.volume_code: old_vol,
        new_vol.volume_code: new_vol,
    }
    old_codes = {
        str((it.get("old") or {}).get("volume_code") or "").strip()
        for it in items
        if it.get("old")
    }
    old_codes.discard("")
    missing = [c for c in old_codes if c not in volumes_by_code]
    if missing:
        for v in Volume.query.filter(Volume.volume_code.in_(missing)).all():
            volumes_by_code[v.volume_code] = v
    return volumes_by_code


def _unique_sheet_name(name: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", (name or "").strip())[:31] or "粗分"
    candidate = base
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = f"{base[: max(1, 31 - len(suffix))]}{suffix}"
        i += 1
    used.add(candidate)
    return candidate


def _write_coarse_sheet(
    wb: openpyxl.Workbook,
    *,
    sheet_name: str,
    title: str,
    meta: str,
    items: list[dict[str, Any]],
    old_vol: Volume,
    new_vol: Volume,
    volumes_by_code: dict[str, Volume],
    used_names: set[str],
) -> None:
    name = _unique_sheet_name(sheet_name, used_names)
    if wb.sheetnames == ["Sheet"]:
        ws = wb.active
        ws.title = name
    else:
        ws = wb.create_sheet(name)

    new_meta = _vol_fields(new_vol)
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    ws["A1"].font = _TITLE_FONT

    ws.append([meta])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_HEADERS))
    ws["A2"].font = _META_FONT

    ws.append(_HEADERS)
    for col, _ in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=3, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _CELL_BORDER

    for idx, it in enumerate(items, start=1):
        new_b = it.get("new") or {}
        old_b = it.get("old")
        old_v = _resolve_old_volume(
            brief=old_b if isinstance(old_b, dict) else None,
            fallback=old_vol if old_b else None,
            volumes_by_code=volumes_by_code,
        )
        old_meta = _vol_fields(old_v) if old_b else {"publisher": "", "grade": "", "volume": ""}
        publisher = new_meta["publisher"] or old_meta["publisher"]
        method = _MATCH_METHOD_LABEL.get(
            str(it.get("match_method") or ""),
            str(it.get("match_method") or "—") or "—",
        )
        row = [
            idx,
            publisher,
            _lesson_name(new_b if isinstance(new_b, dict) else None),
            new_meta["grade"],
            new_meta["volume"],
            _lesson_unit(new_b if isinstance(new_b, dict) else None),
            _lesson_hour(new_b if isinstance(new_b, dict) else None),
            method,
            _lesson_name(old_b if isinstance(old_b, dict) else None) or "—",
            old_meta["grade"],
            old_meta["volume"],
            _lesson_unit(old_b if isinstance(old_b, dict) else None),
            _lesson_hour(old_b if isinstance(old_b, dict) else None),
            _lesson_courseware_id(old_b if isinstance(old_b, dict) else None),
        ]
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(_HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _CELL_BORDER
            cell.alignment = _CENTER if c in _CENTER_COLS else _WRAP

    widths = [6, 12, 22, 10, 8, 22, 22, 12, 22, 10, 8, 22, 22, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[3].height = 20


def _load_pair_items(old_code: str, new_code: str) -> tuple[Volume, Volume, list[dict[str, Any]]]:
    old_vol = get_diff_volume_by_code(old_code)
    new_vol = get_diff_volume_by_code(new_code)
    coarse = list_stored_pairs(
        old_vol_id=old_vol.id,
        new_vol_id=new_vol.id,
        old_code=old_code,
        new_code=new_code,
        enrich_changes=False,
    )
    return old_vol, new_vol, list(coarse.get("items") or [])


def build_coarse_match_export_xlsx(
    *,
    old_code: str,
    new_code: str,
    include_matched: bool = True,
    include_unmatched: bool = True,
) -> tuple[bytes, str]:
    """导出课时粗分结果 Excel。返回 (bytes, filename)。"""
    old_code = (old_code or "").strip()
    new_code = (new_code or "").strip()
    if not old_code or not new_code:
        raise ValueError("缺少 old_code / new_code")
    if not include_matched and not include_unmatched:
        raise ValueError("请至少勾选「已匹配」或「未匹配」")

    old_vol, new_vol, items = _load_pair_items(old_code, new_code)
    items = _filter_items(
        items, include_matched=include_matched, include_unmatched=include_unmatched
    )
    if not items:
        raise ValueError("尚无课时粗分结果，请先运行粗分")

    volumes_by_code = _collect_volumes_by_code(items, old_vol=old_vol, new_vol=new_vol)
    new_meta = _vol_fields(new_vol)
    publisher_name = new_meta["publisher"] or str(old_vol.edition or "").strip() or "粗分结果"
    sheet_name = _export_sheet_name(new_meta["grade"], new_meta["volume"])

    wb = openpyxl.Workbook()
    _write_coarse_sheet(
        wb,
        sheet_name=sheet_name,
        title=f"{new_vol.display_title or new_code} · 课时粗分结果",
        meta=(
            f"新册 {new_code}  ↔  旧册 {old_code}  ·  "
            f"{new_vol.edition or ''}{new_vol.subject or ''}"
            f"{_grade_label(new_vol.grade)}{_semester_label(new_vol.semester)}  ·  "
            f"共 {len(items)} 条"
        ),
        items=items,
        old_vol=old_vol,
        new_vol=new_vol,
        volumes_by_code=volumes_by_code,
        used_names=set(),
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), _export_filename(publisher_name)


def list_xiaoke_edition_coarse_export_targets(edition_id: str) -> dict[str, Any]:
    """小科版本页：各册粗分导出可选状态。"""
    from ...models import DiffLessonPair
    from ..old_library.edition_registry import get_edition, grade_term_pairs_for_edition
    from .sandbox_subjects import xiaoke_diff_prefix
    from .volume_create import make_diff_volume_code

    edition = get_edition(edition_id)
    prefix = xiaoke_diff_prefix(edition)
    volumes: list[dict[str, Any]] = []
    for grade, term in grade_term_pairs_for_edition(edition):
        old_code = make_diff_volume_code(
            grade=grade, term=term, book_type="diff_old", prefix=prefix
        )
        new_code = make_diff_volume_code(
            grade=grade, term=term, book_type="diff_new", prefix=prefix
        )
        label = f"{_grade_label(grade)}{_semester_label(term)}"
        try:
            old_vol = get_diff_volume_by_code(old_code)
            new_vol = get_diff_volume_by_code(new_code)
        except ValueError:
            volumes.append(
                {
                    "grade": grade,
                    "term": term,
                    "label": label,
                    "old_code": old_code,
                    "new_code": new_code,
                    "has_coarse": False,
                    "matched_count": 0,
                    "unmatched_count": 0,
                    "pair_count": 0,
                }
            )
            continue
        rows = DiffLessonPair.query.filter_by(
            old_volume_id=old_vol.id, new_volume_id=new_vol.id
        ).all()
        matched_n = 0
        unmatched_n = 0
        for r in rows:
            method = str(r.match_method or "").strip().lower()
            if method in ("none", "unmatched", "no_match") or not r.old_lesson_id:
                unmatched_n += 1
            else:
                matched_n += 1
        volumes.append(
            {
                "grade": grade,
                "term": term,
                "label": label,
                "old_code": old_code,
                "new_code": new_code,
                "has_coarse": len(rows) > 0,
                "matched_count": matched_n,
                "unmatched_count": unmatched_n,
                "pair_count": len(rows),
            }
        )
    return {
        "edition_id": edition.edition_id,
        "edition_label": edition.label,
        "volumes": volumes,
    }


def build_edition_coarse_match_export_xlsx(
    *,
    edition_id: str,
    pairs: list[dict[str, str]] | None = None,
    include_matched: bool = True,
    include_unmatched: bool = True,
) -> tuple[bytes, str]:
    """整版导出：文件名=版本名.xlsx，每册一个 sheet（四年级上册…）。"""
    if not include_matched and not include_unmatched:
        raise ValueError("请至少勾选「已匹配」或「未匹配」")

    from ..old_library.edition_registry import get_edition

    edition = get_edition(edition_id)
    targets = list_xiaoke_edition_coarse_export_targets(edition.edition_id)
    wanted: set[tuple[str, str]] | None = None
    if pairs:
        wanted = set()
        for p in pairs:
            oc = str((p or {}).get("old_code") or "").strip()
            nc = str((p or {}).get("new_code") or "").strip()
            if oc and nc:
                wanted.add((oc, nc))
        if not wanted:
            raise ValueError("请至少勾选一册")

    selected = []
    for vol in targets["volumes"]:
        key = (vol["old_code"], vol["new_code"])
        if wanted is not None and key not in wanted:
            continue
        if not vol["has_coarse"]:
            continue
        selected.append(vol)
    if not selected:
        raise ValueError("所选册次尚无粗分结果，请先在本册建设运行粗分")

    wb = openpyxl.Workbook()
    # 去掉默认空表，由首册写入
    default = wb.active
    used: set[str] = set()
    total_rows = 0
    for vol in selected:
        old_vol, new_vol, items = _load_pair_items(vol["old_code"], vol["new_code"])
        items = _filter_items(
            items, include_matched=include_matched, include_unmatched=include_unmatched
        )
        if not items:
            continue
        volumes_by_code = _collect_volumes_by_code(
            items, old_vol=old_vol, new_vol=new_vol
        )
        _write_coarse_sheet(
            wb,
            sheet_name=vol["label"],
            title=f"{new_vol.display_title or vol['new_code']} · 课时粗分结果",
            meta=(
                f"新册 {vol['new_code']}  ↔  旧册 {vol['old_code']}  ·  "
                f"{edition.label}  ·  共 {len(items)} 条"
            ),
            items=items,
            old_vol=old_vol,
            new_vol=new_vol,
            volumes_by_code=volumes_by_code,
            used_names=used,
        )
        total_rows += len(items)

    if total_rows <= 0:
        raise ValueError("按当前筛选无课时可导出（请勾选已匹配/未匹配，或换有结果的册）")

    # 若首张仍是空 Sheet（理论上不会），删掉
    if default.title == "Sheet" and default.max_row <= 1 and len(wb.sheetnames) > 1:
        wb.remove(default)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), _export_filename(edition.label)
